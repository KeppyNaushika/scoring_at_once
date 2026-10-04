"""Excel (openpyxl) を書き出すモジュールで共通に使う書式と部品."""

from __future__ import annotations

from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

FONT = Font(size=11, name="Meiryo UI")
CENTER = Alignment(horizontal="center", vertical="center")


def writable_cell(sheet: Worksheet, row: int, column: int) -> Cell:
    """値を書き込むセル. 結合されたセルの左上以外 (MergedCell) は値を持てないので, 書き込む前に確かめる."""
    cell = sheet.cell(row=row, column=column)
    assert isinstance(cell, Cell), f"結合されたセルには書き込めません: {cell.coordinate}"
    return cell


def solid_fill(color: str) -> PatternFill:
    return PatternFill(patternType="solid", fgColor=color)


def center_all(sheet: Worksheet) -> None:
    """値を書いた範囲の全てのセルを既定のフォント・中央揃えにする."""
    for cell in (cell for row in sheet.rows for cell in row):
        cell.font = FONT
        cell.alignment = CENTER
