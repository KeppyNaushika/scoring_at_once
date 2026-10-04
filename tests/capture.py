"""アプリを外側から操作して, 結果 (保存された JSON・書き出したファイル) を JSON にまとめる.

画面の作りに依存しないよう, ボタンは表示文字で探し, 操作はキー入力やクリックのイベントで行う.
作り直す前後で同じ操作をして結果を比べることで, 挙動が変わっていないことを確かめる.

使い方: python tests/capture.py <アプリのディレクトリ> <シナリオ名> <出力 JSON>
"""

from __future__ import annotations

import hashlib
import json
import os
import shutil
import sqlite3
import sys
import tempfile
import traceback
import zipfile
from pathlib import Path
from typing import Any, Callable, Iterator

APP_DIR, SCENARIO, OUTPUT = sys.argv[1], sys.argv[2], Path(sys.argv[3])
sys.path.insert(0, str(Path(__file__).parent))
sys.path.insert(0, APP_DIR)

import fixture  # noqa: E402

EXPORT_VARIANTS: dict[str, dict[str, Any]] = {
    "default": fixture.EXPORT_DEFAULT,
    "corners": {
        "symbol": {"position": "nw", "x": 5, "y": -3, "size": 40, "unscored": False, "correct": True, "partial": True, "hold": False, "incorrect": True},
        "point": {"position": "se", "x": -10, "y": 4, "size": 20, "unscored": True, "correct": False, "partial": True, "hold": True, "incorrect": False},
    },
    "edges": {
        "symbol": {"position": "e", "x": 0, "y": 10, "size": 80, "unscored": True, "correct": True, "partial": False, "hold": True, "incorrect": True},
        "point": {"position": "n", "x": 3, "y": 0, "size": 12, "unscored": False, "correct": True, "partial": True, "hold": False, "incorrect": True},
    },
}

if SCENARIO == "sao":
    # .sao の ID は試験フォルダの絶対パスから作られるので, 毎回同じ場所を使う
    tmp = Path(tempfile.gettempdir()) / "scoring_at_once_capture_sao"
    shutil.rmtree(tmp, ignore_errors=True)
    tmp.mkdir()
else:
    tmp = Path(tempfile.mkdtemp(prefix="scoring_at_once_capture_"))
variant = SCENARIO.split(":")[1] if ":" in SCENARIO else "default"
paths = fixture.build(tmp, EXPORT_VARIANTS.get(variant, fixture.EXPORT_DEFAULT))
os.environ["SCORING_AT_ONCE_CONFIG_DIR"] = str(paths["config_dir"])

import subprocess  # noqa: E402
import tkinter  # noqa: E402
import tkinter.filedialog  # noqa: E402
import tkinter.messagebox  # noqa: E402
import tkinter.simpledialog  # noqa: E402

# ダイアログと外部アプリの起動を止める
messages: list[str] = []
for name in ("showinfo", "showwarning", "showerror"):
    setattr(tkinter.messagebox, name, lambda title="", message="", **k: messages.append(str(title)))
for name in ("askokcancel", "askyesno"):
    setattr(tkinter.messagebox, name, lambda *a, **k: True)
setattr(tkinter.simpledialog, "askstring", lambda *a, **k: "tester")
setattr(tkinter.filedialog, "asksaveasfilename", lambda **k: str(tmp / f"out.{k.get('defaultextension', 'bin').lstrip('.')}"))
setattr(subprocess, "run", lambda *a, **k: None)
if hasattr(os, "startfile"):
    setattr(os, "startfile", lambda *a, **k: None)

errors: list[str] = []
setattr(tkinter.Tk, "report_callback_exception", lambda self, *a: errors.append(
    "".join(traceback.format_exception(*a)).strip().splitlines()[-1]
))

import score  # noqa: E402

score.main = lambda: None  # 子ウインドウを閉じたときのアプリ再起動を止める

result: dict[str, Any] = {}


def widgets(root: tkinter.Misc, kind: type) -> Iterator[Any]:
    for child in root.winfo_children():
        if isinstance(child, kind):
            yield child
        yield from widgets(child, kind)


def button(root: tkinter.Misc, text: str, nth: int = 0) -> tkinter.Button:
    found = [b for b in widgets(root, tkinter.Button) if b.cget("text") == text]
    if len(found) <= nth:
        raise LookupError(f"ボタン {text!r} がありません: {[b.cget('text') for b in widgets(root, tkinter.Button)]}")
    return found[nth]


def toplevel() -> tkinter.Toplevel:
    tops = [w for w in widgets(app_root, tkinter.Toplevel) if w.winfo_exists()]
    return tops[-1]


def press(window: tkinter.Misc, *keys: str) -> None:
    window.focus_force()
    for key in keys:
        window.event_generate(key if key.startswith("<") else f"<KeyPress-{key}>")
        window.update()


def file_md5(path: Path) -> str:
    return hashlib.md5(path.read_bytes()).hexdigest()


def read_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def dump_xlsx(path: Path) -> list[Any]:
    import openpyxl

    workbook = openpyxl.load_workbook(path)
    rows: list[Any] = []
    for sheet in workbook.worksheets:
        rows.append([sheet.title, "merged", sorted(str(r) for r in sheet.merged_cells.ranges)])
        rows.append([sheet.title, "widths", sorted((k, v.width) for k, v in sheet.column_dimensions.items())])
        rows.append([sheet.title, "images", sorted((i.anchor._from.col, i.anchor._from.row, i.width, i.height) for i in getattr(sheet, "_images"))])
        for row in sheet.iter_rows():
            for c in row:
                if c.value is None and c.fill.fgColor.rgb == "00000000":
                    continue
                rows.append([sheet.title, c.coordinate, c.value, c.fill.fgColor.rgb, c.protection.locked, c.number_format])
    return rows


def finish() -> None:
    result["errors"] = errors
    result["messages"] = messages
    OUTPUT.write_text(json.dumps(result, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
    shutil.rmtree(tmp, ignore_errors=True)
    os._exit(0)


def step(delay: int, action: Callable[[], None]) -> None:
    def run() -> None:
        try:
            action()
        except Exception:
            result["exception"] = traceback.format_exc().strip().splitlines()[-1]
            finish()
    app_root.after(delay, run)


# ---------------------------------------------------------------------------
# シナリオ
# ---------------------------------------------------------------------------
def scenario_roster() -> None:
    """名簿/配点の Excel を作り, 値を入れて読み込む."""
    def make() -> None:
        button(app_root, "名簿/配点を\nExcel で入力").invoke()
        xlsx = paths["work_dir"] / "名簿と配点の入力.xlsx"
        result["roster_xlsx"] = dump_xlsx(xlsx)
        import openpyxl

        workbook = openpyxl.load_workbook(xlsx)
        sheet = workbook["名簿登録"]
        sheet["B2"] = "3"
        sheet["F3"] = "新しい 名前"
        sheet["E7"] = 12345
        haiten = workbook["配点登録"]
        haiten["C5"] = 9
        haiten["F6"] = None
        haiten["F7"] = 7
        haiten["E4"] = "イ"
        workbook.save(xlsx)
        button(app_root, "名簿/配点を\n読み込む").invoke()
        result["meibo"] = read_json(paths["work_dir"] / "meibo.json")
        result["questions"] = [
            {k: v for k, v in q.items() if k != "score"}
            for q in read_json(paths["work_dir"] / "answer_area.json")["questions"]
        ]
        result["xlsx_left"] = xlsx.exists()
        finish()
    step(10, make)


def scenario_export() -> None:
    """書き出し画面で PDF と採点結果一覧 Excel を出力する."""
    step(10, lambda: button(app_root, "書き出す").invoke())

    def run() -> None:
        window = toplevel()
        window.update()
        button(window, "採点済答案画像の出力").invoke()
        window.update()
        button(window, "採点結果一覧表(.xlsx)の出力").invoke()
        window.update()
        result["png"] = {p.name: file_md5(p) for p in sorted((paths["work_dir"] / "output").glob("*.png"))}
        result["pdf_exists"] = (tmp / "out.pdf").exists()
        result["result_xlsx"] = dump_xlsx(tmp / "out.xlsx")
        finish()
    step(2500, run)


def scenario_preset(name: str) -> Callable[[], None]:
    def scenario() -> None:
        step(10, lambda: button(app_root, "書き出す").invoke())

        def run() -> None:
            window = toplevel()
            window.update()
            button(window, name).invoke()
            window.update()
            result["export_settings"] = read_json(paths["config_dir"] / "config.json")["projects"][0]["export"]
            button(window, "採点済答案画像の出力").invoke()
            result["png"] = {p.name: file_md5(p) for p in sorted((paths["work_dir"] / "output").glob("*.png"))}
            finish()
        step(2500, run)
    return scenario


def scenario_grading() -> None:
    """一括採点画面でキー操作して採点する."""
    step(10, lambda: button(app_root, "一括採点する").invoke())

    def run() -> None:
        window = toplevel()
        window.update()
        # 全ての採点状態を表示する (初期状態は未採点のみ)
        press(window, "<Control-e>", "<Control-f>", "<Control-j>", "<Control-o>")
        press(window, "e", "2", "f", "o", "d", "s", "a", "w", "3", "<BackSpace>", "4", "j", "q", "s", "s", "e")
        listbox = next(widgets(window, tkinter.Listbox))
        listbox.selection_clear(0, "end")
        listbox.selection_set(2)
        listbox.event_generate("<<ListboxSelect>>")
        window.update()
        press(window, "o", "e", "1", "f", "d", "d", "j", "r")
        result["scores"] = [
            q["score"] for q in read_json(paths["work_dir"] / "answer_area.json")["questions"]
        ]
        finish()
    step(2500, run)


def scenario_area() -> None:
    """解答欄の指定画面で枠を作り, 種類・順番を変え, 削除する."""
    step(10, lambda: button(app_root, "解答欄の位置を指定").invoke())

    def run() -> None:
        window = toplevel()
        window.update()
        canvas = max(widgets(window, tkinter.Canvas), key=lambda c: c.winfo_width() * c.winfo_height())
        canvas.event_generate("<ButtonPress-1>", x=100, y=550)
        canvas.event_generate("<B1-Motion>", x=180, y=570)
        canvas.event_generate("<B1-Motion>", x=230, y=590)
        canvas.event_generate("<ButtonRelease-1>", x=230, y=590)
        window.update()
        button(window, "小計点").invoke()
        button(window, "上へ").invoke()
        button(window, "上へ").invoke()
        window.update()
        # 逆向きにドラッグしても左上が (x0, y0) になる
        canvas.event_generate("<ButtonPress-1>", x=300, y=300)
        canvas.event_generate("<B1-Motion>", x=250, y=260)
        canvas.event_generate("<ButtonRelease-1>", x=250, y=260)
        window.update()
        button(window, "氏名").invoke()
        listbox = next(widgets(window, tkinter.Listbox))
        listbox.selection_clear(0, "end")
        listbox.selection_set(3)
        listbox.event_generate("<<ListboxSelect>>")
        window.update()
        button(window, "下へ").invoke()
        button(window, "削除").invoke()
        window.update()
        data = read_json(paths["work_dir"] / "answer_area.json")["questions"]
        result["regions"] = [
            {k: v for k, v in q.items() if k != "score"} | {"n_scores": len(q["score"])} for q in data
        ]
        finish()
    step(2500, run)


def scenario_projects() -> None:
    """試験一覧の並べ替えと削除."""
    config_path = paths["config_dir"] / "config.json"
    config = read_json(config_path)
    for name in ("二つ目", "三つ目"):
        config["projects"].append(dict(config["projects"][0], name=name))
    config_path.write_text(json.dumps(config, ensure_ascii=False))

    def run() -> None:
        listbox = next(widgets(app_root, tkinter.Listbox))
        listbox.selection_clear(0, "end")
        listbox.selection_set(1)
        listbox.event_generate("<<ListboxSelect>>")
        app_root.update()
        names = lambda: [p["name"] for p in read_json(config_path)["projects"]]  # noqa: E731
        button(app_root, "上へ").invoke()
        result["after_up"] = names()
        button(app_root, "上へ").invoke()
        result["after_up_top"] = names()
        button(app_root, "下へ").invoke()
        button(app_root, "下へ").invoke()
        button(app_root, "下へ").invoke()
        result["after_down"] = names()
        result["listbox"] = list(listbox.get(0, "end"))
        button(app_root, "削除").invoke()
        result["after_delete"] = names()
        result["selected"] = read_json(config_path)["index_projects_in_listbox"]
        finish()
    step(500, run)


def scenario_sao() -> None:
    """後継版へ書き出す."""
    def run() -> None:
        button(app_root, "後継版へ書き出す\n(.sao)").invoke()
        sao = tmp / "out.sao"
        with zipfile.ZipFile(sao) as archive:
            names = sorted(archive.namelist())
            manifest = json.loads(archive.read("manifest.json"))
            db_path = tmp / "archive.db"
            db_path.write_bytes(archive.read("archive.db"))
            files = {n: hashlib.md5(archive.read(n)).hexdigest() for n in names if n.startswith("files/")}
        for key in ("exportedAt", "appVersion"):
            manifest.pop(key)
        db = sqlite3.connect(db_path)
        rows = {}
        for table in manifest["rowCounts"]:
            cursor = db.execute(f'SELECT * FROM "{table}" ORDER BY id')
            columns = [d[0] for d in cursor.description]
            rows[table] = [
                {c: v for c, v in zip(columns, r) if c not in ("createdAt", "updatedAt", "invitedAt", "referenceDate", "startDate")}
                for r in cursor
            ]
        result.update(names=names, manifest=manifest, rows=rows, files=files)
        finish()
    step(10, run)


SCENARIOS: dict[str, Callable[[], None]] = {
    "roster": scenario_roster,
    "grading": scenario_grading,
    "area": scenario_area,
    "projects": scenario_projects,
    "sao": scenario_sao,
    "preset1": scenario_preset("例1"),
    "preset2": scenario_preset("例2"),
}

app_root = tkinter.Tk()
app_root.geometry("800x500")
score.MainFrame(root=app_root)
SCENARIOS.get(SCENARIO.split(":")[0], scenario_export)()
def timeout() -> None:
    result.setdefault("exception", "timeout")
    finish()


app_root.after(60000, timeout)
app_root.mainloop()
